"""Temporary probe round 3: House consolidated CPF xlsx, archived CPF page, House gift travel search, Senate eFD report types."""
import io, json, re, sys, urllib.request, urllib.error, urllib.parse
from pathlib import Path
sys.path.insert(0, "tools")
OUT = Path("tools/probe/out"); OUT.mkdir(parents=True, exist_ok=True)
UA = {"User-Agent": "ThePillory/1.0 (+https://thepillory.co; public records)", "Accept": "*/*"}
def get(url, data=None, headers=None, n=20000000):
    h = dict(UA); h.update(headers or {})
    try:
        with urllib.request.urlopen(urllib.request.Request(url, data=data, headers=h), timeout=120) as r:
            return r.status, r.headers.get("Content-Type"), r.read()[:n]
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Content-Type"), e.read()[:3000]
    except Exception as e:
        return 0, None, repr(e).encode()
def links(b, pat):
    return sorted(set(re.findall(r'href="([^"]*(?:%s)[^"]*)"' % pat, b.decode("utf-8", "replace"), re.I)))[:120]
res = {}
s, ct, b = get("https://appropriations.house.gov/sites/evo-subsites/republicans-appropriations.house.gov/files/evo-media-document/fy26-house-cpf-consolidated.xlsx")
res["xlsx"] = [s, ct, len(b)]
try:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True)
    res["sheets"] = wb.sheetnames
    for ws in wb.worksheets[:3]:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append([str(c)[:60] if c is not None else None for c in row])
            if i > 12: break
        res[f"sheet_{ws.title}"] = {"dims": ws.max_row, "rows": rows}
except Exception as e:
    res["xlsx_err"] = repr(e)
s, ct, b = get("https://appropriations.house.gov/committee-activity/archived-community-project-funding")
res["archived"] = links(b, r"\.xlsx|\.xls|\.csv|enact|cpf|communit")
s, ct, b = get("https://appropriations.house.gov/fy27-information/fy27-community-project-funding")
res["fy27"] = links(b, r"\.xlsx|\.xls|\.csv|\.pdf|request|enact")
# House gift travel search
for url, data in [("https://disclosures-clerk.house.gov/GiftTravelFilings/ViewSearch", None),
                  ("https://disclosures-clerk.house.gov/GiftTravelFilings/ViewSearch", urllib.parse.urlencode({"FilingYear": "2025"}).encode())]:
    s, ct, b = get(url, data=data, headers={"X-Requested-With": "XMLHttpRequest"} if data else None)
    h = b.decode("utf-8", "replace")
    res[f"gt_{'post' if data else 'get'}"] = {"status": s, "ct": ct, "inputs": re.findall(r'<(?:input|select)[^>]*name="([^"]+)"', h)[:40], "options": re.findall(r'<option[^>]*value="([^"]*)"', h)[:40],
        "forms": re.findall(r'<form[^>]*>', h)[:5], "links": links(b, r"pdf|travel|gift|zip|download"), "text": re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", h))[:1500]}
for y in (2025,):
    for path in (f"public_disc/gift-travel-pdfs/{y}.zip", f"public_disc/financial-pdfs/{y}FD.zip", f"public_disc/gifttravel/{y}GT.zip", f"public_disc/gift-travel/{y}.zip"):
        s, ct, b = get(f"https://disclosures-clerk.house.gov/{path}", n=300)
        res[f"z_{path}"] = [s, ct]
# Senate eFD: the search page's report types after accepting the terms
from load_trades import Senate
sen = Senate(); sen.start()
page = sen._req("https://efdsearch.senate.gov/search/")
res["efd_types"] = re.findall(r'<input[^>]*name="report_type"[^>]*value="(\d+)"[^>]*>\s*([^<]*)', page)[:30] or re.findall(r'value="(\d+)"[^>]*>\s*<[^>]*>?\s*([A-Za-z][^<]{2,60})', page)[:40]
res["efd_text"] = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", page))[:2500]
(OUT / "govmoney3.json").write_text(json.dumps(res, indent=1, default=str))
